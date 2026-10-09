"""Shared parsing, validation, and surgical XLSX editing for GLOMBO.

Python 3.10+, openpyxl (read only), lxml. The OOXML writer changes cell values
inside existing sheets and copies every other XLSX package part unchanged.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import posixpath
import re
import zipfile
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree
from openpyxl import load_workbook

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
Q = lambda name: f'{{{NS}}}{name}'
CELL = re.compile(r'^([A-Z]+)([1-9][0-9]*)$')


class InputError(ValueError):
    pass


def utc(value: str) -> dt.datetime:
    if not value or not value.endswith('Z'):
        raise InputError(f'UTC timestamp must end in Z: {value!r}')
    try:
        parsed = dt.datetime.fromisoformat(value.replace('Z', '+00:00'))
        if parsed.utcoffset() != dt.timedelta(0):
            raise ValueError()
        return parsed
    except ValueError as exc:
        raise InputError(f'Invalid UTC timestamp: {value!r}') from exc


def decimal(value: str, low: float, high: float, label: str):
    if value == '':
        return None
    try:
        number = float(value)
    except (ValueError, TypeError) as exc:
        raise InputError(f'{label} must be numeric: {value!r}') from exc
    if not math.isfinite(number) or not low <= number <= high:
        raise InputError(f'{label} outside {low}..{high}: {value!r}')
    return number


def read_samples(path: Path, allow_mock=False):
    with path.open(newline='', encoding='utf-8-sig') as file:
        reader = csv.DictReader(file)
        required = {'record_status','sample_event_id', 'filter_id', 'filter_position',
                    'sample_start_utc', 'sample_end_utc', 'sampled_volume_ml',
                    'error_code', 'sampling_status', 'kit_uid', 'environment_type'}
        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise InputError('Standard CSV missing columns: ' + ', '.join(sorted(required-set(reader.fieldnames or []))))
        rows = list(reader)
    if not rows:
        raise InputError('Standard CSV has no sample rows')
    if any(not (row.get('record_status') or '').strip() for row in rows):
        raise InputError('record_status must identify production or test data on every row')
    if any('MOCK' in row['record_status'].upper() for row in rows) and not allow_mock:
        raise InputError('Mock records require --allow-mock; never use test data for submission')
    ids = [(row.get('sample_event_id') or '').strip() for row in rows]
    if any(not x for x in ids) or len(set(ids)) != len(ids):
        raise InputError('sample_event_id must be nonempty and unique, including failed attempts')
    kits = [r['kit_uid'].strip() for r in rows if r['kit_uid'].strip()]
    if len(kits) != len(set(kits)):
        raise InputError('A physical kit_uid appears in more than one event')
    for i, row in enumerate(rows, 2):
        if not (row.get('filter_position') or '').strip():
            raise InputError(f'Row {i}: filter_position missing')
        start, end = utc(row['sample_start_utc']), utc(row['sample_end_utc'])
        if end < start:
            raise InputError(f'Row {i}: sample end precedes start')
        row['_start'] = start
        row['_end'] = end
        row['_volume'] = decimal(row['sampled_volume_ml'], 0, 1e8, 'sampled_volume_ml')
        row['_lat'] = decimal(row.get('latitude',''), -90, 90, 'latitude')
        row['_lon'] = decimal(row.get('longitude',''), -180, 180, 'longitude')
        if (row['_lat'] is None) != (row['_lon'] is None):
            raise InputError(f'Row {i}: latitude and longitude must be paired')
        if row.get('coordinate_system','') and row['coordinate_system'] != 'WGS84':
            raise InputError(f'Row {i}: only confirmed WGS84 decimal degrees supported')
        if row['_lat'] is not None and row.get('coordinate_system','') != 'WGS84':
            raise InputError(f'Row {i}: coordinate_system=WGS84 required when coordinates provided')
        if row['sampling_status'] not in {'complete','failed','aborted'}:
            raise InputError(f'Row {i}: unknown sampling_status')
        category=row.get('sample_category','sample') or 'sample'
        if category not in {'sample','negative control','positive control','PCR standard'}:
            raise InputError(f'Row {i}: sample_category outside supported FAIRe values')
        if category=='negative control' and not row.get('neg_cont_type'):
            raise InputError(f'Row {i}: negative control requires neg_cont_type')
        if category=='positive control' and not row.get('pos_cont_type'):
            raise InputError(f'Row {i}: positive control requires pos_cont_type')
        try:
            row['_error'] = int(row['error_code'])
        except ValueError as exc:
            raise InputError(f'Row {i}: error_code must be an integer') from exc
        if row['sampling_status']=='complete' and row['_error']!=0:
            raise InputError(f'Row {i}: complete sampling has nonzero error')
        if row['sampling_status']!='complete' and row['_error']==0:
            raise InputError(f'Row {i}: incomplete sampling has zero error')
    return rows


def samples_for_output(rows):
    return [r for r in rows if r['sampling_status']=='complete' and r['_error']==0]


def read_json(path: Path):
    data = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(data, dict):
        raise InputError('Configuration must be a JSON object')
    return data


def sheet_headers(workbook: Path, sheet: str, row: int):
    wb = load_workbook(workbook, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise InputError(f'Workbook lacks sheet {sheet}')
    values = next(wb[sheet].iter_rows(min_row=row, max_row=row, values_only=True))
    wb.close()
    return {str(value): i for i, value in enumerate(values, 1) if value is not None}


def col(n: int) -> str:
    s = ''
    while n:
        n, rem = divmod(n-1,26)
        s = chr(65+rem)+s
    return s


def cell(colnum: int, rownum: int):
    return f'{col(colnum)}{rownum}'


def report_csv(path: Path, headers, rows):
    with path.open('w', newline='', encoding='utf-8') as file:
        writer=csv.DictWriter(file,fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def _sheet_paths(source: zipfile.ZipFile):
    book=etree.fromstring(source.read('xl/workbook.xml'))
    rels=etree.fromstring(source.read('xl/_rels/workbook.xml.rels'))
    targets={r.get('Id'):r.get('Target') for r in rels.findall(f'{{{PKG}}}Relationship')}
    answer={}
    for s in book.find(Q('sheets')):
        target=targets[s.get(f'{{{REL}}}id')]
        answer[s.get('name')]=posixpath.normpath(posixpath.join('xl',target.lstrip('/'))) if not target.startswith('/xl/') else target.lstrip('/')
    return answer


def _column_number(address):
    match=CELL.fullmatch(address)
    if not match: raise InputError(f'Invalid cell address {address}')
    n=0
    for letter in match[1]: n=n*26+ord(letter)-64
    return n,int(match[2])


def _modify_sheet(xml: bytes, edits: dict):
    root=etree.fromstring(xml, parser=etree.XMLParser(remove_blank_text=False))
    data=root.find(Q('sheetData'))
    rows={int(row.get('r')):row for row in data.findall(Q('row'))}
    max_existing=max(rows,default=0)
    for address, value in edits.items():
        ci,ri=_column_number(address)
        row=rows.get(ri)
        if row is None:
            row=etree.Element(Q('row'),r=str(ri))
            if ri>max_existing:
                data.append(row)
                max_existing=ri
            else:
                following=next((x for x in data if int(x.get('r'))>ri),None)
                if following is None:data.append(row)
                else:data.insert(data.index(following),row)
            rows[ri]=row
        existing=next((x for x in row if x.get('r')==address),None)
        if existing is None:
            existing=etree.Element(Q('c'),r=address)
            following=next((x for x in row if _column_number(x.get('r'))[0]>ci),None)
            if following is None:row.append(existing)
            else:row.insert(row.index(following),existing)
        for child in list(existing):existing.remove(child)
        existing.attrib.pop('t',None)
        if value is None or value=='':continue
        if isinstance(value,(float,int)) and not isinstance(value,bool):
            etree.SubElement(existing,Q('v')).text=str(value)
        else:
            existing.set('t','inlineStr')
            ist=etree.SubElement(existing,Q('is'))
            t=etree.SubElement(ist,Q('t'))
            t.text=str(value)
            if str(value).strip()!=str(value):t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    # The template's worksheet dimension may describe only its original three header rows.
    dim=root.find(Q('dimension'))
    if dim is not None:
        max_row=max(rows) if rows else 1
        max_col=max((_column_number(c.get('r'))[0] for row in rows.values() for c in row if c.tag==Q('c')),default=1)
        dim.set('ref',f'A1:{col(max_col)}{max_row}')
    return etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)


def patch_workbook(original: Path, target: Path, changes: dict[str,dict]):
    if original.resolve()==target.resolve():
        raise InputError('Output must be a new file; never overwrite the supplied template')
    target.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(original) as source:
        paths=_sheet_paths(source)
        missing=set(changes)-set(paths)
        if missing:raise InputError(f'Template sheets missing: {missing}')
        by_path={paths[name]:updates for name,updates in changes.items()}
        with zipfile.ZipFile(target,'w') as output:
            for member in source.infolist():
                raw=source.read(member.filename)
                if member.filename in by_path:
                    raw=_modify_sheet(raw,by_path[member.filename])
                output.writestr(member,raw)
