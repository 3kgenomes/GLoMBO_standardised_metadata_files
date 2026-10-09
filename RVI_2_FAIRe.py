#!/usr/bin/env python3
"""Fill the supplied FAIRe v1.0.2 FULL template with supported sampler metadata.

This is a publication preparation draft. Other FAIRe components and fields
remain in the original template for later validated laboratory/project data.
"""
import argparse
from pathlib import Path

from openpyxl import load_workbook
from glombo_common import (InputError, cell, patch_workbook, read_json,
                           read_samples, report_csv, samples_for_output,
                           sheet_headers)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--samples',type=Path,required=True)
    p.add_argument('--template',type=Path,required=True)
    p.add_argument('--project',type=Path,required=True,help='Confirmed project metadata JSON')
    p.add_argument('--out',type=Path,required=True)
    p.add_argument('--allow-mock',action='store_true',help='Only for explicit test fixtures')
    args=p.parse_args()
    source=read_samples(args.samples,allow_mock=args.allow_mock)
    rows=samples_for_output(source)
    if not rows:raise InputError('No completed zero-error sampling events')
    config=read_json(args.project)
    wb=load_workbook(args.template,read_only=True,data_only=True)
    if any(wb[name].max_row != 3 for name in
           ('sampleMetadata','ampData','stdData','eLowQuantData',
            'experimentRunMetadata','taxaRaw','taxaFinal')):
        raise InputError('Start from the unpopulated FULL template; no existing data will be replaced')
    terms={str(row[2]):i for i,row in enumerate(wb['projectMetadata'].values,1) if row[2]}
    wb.close()
    sample_terms=sheet_headers(args.template,'sampleMetadata',3)
    changes={'projectMetadata':{},'sampleMetadata':{}}
    project_fields={k:v for k,v in config.items() if k in terms and v not in (None,'')}
    for key,value in project_fields.items():
        changes['projectMetadata'][cell(4,terms[key])]=value
    # Version comes from the user's exact supplied checklist, not another project.
    changes['projectMetadata'][cell(4,terms['checkls_ver'])]='1.0.2'
    selected=['samp_name','samp_category','neg_cont_type','pos_cont_type',
              'materialSampleID','decimalLongitude',
              'decimalLatitude','verbatimCoordinateSystem','verbatimSRS',
              'geo_loc_name','eventDate','env_broad_scale','env_local_scale',
              'env_medium','samp_collect_device','samp_size','samp_size_unit',
              'assay_name',
              'minimumDepthInMeters','maximumDepthInMeters']
    absent=set(selected)-set(sample_terms)
    if absent:raise InputError(f'FAIRe template schema changed: {absent}')
    required_for_this_stage=['decimalLongitude','decimalLatitude','geo_loc_name',
                             'eventDate','env_broad_scale','env_local_scale','env_medium']
    review=[]
    for index,r in enumerate(rows,4):
        fields={
            'samp_name':r['sample_event_id'],
            'samp_category':r.get('sample_category','sample') or 'sample',
            'neg_cont_type':r.get('neg_cont_type') or None,
            'pos_cont_type':r.get('pos_cont_type') or None,
            'materialSampleID':r.get('kit_uid','') or None,
            'decimalLongitude':r['_lon'],'decimalLatitude':r['_lat'],
            'verbatimCoordinateSystem':'decimal degrees' if r['_lat'] is not None else None,
            'verbatimSRS':'WGS84' if r['_lat'] is not None else None,
            'geo_loc_name':r.get('geo_loc_name') or None,
            'eventDate':r['sample_start_utc'],
            'env_broad_scale':r.get('env_broad_scale') or None,
            'env_local_scale':r.get('env_local_scale') or None,
            'env_medium':r.get('env_medium') or None,
            'samp_collect_device':r.get('sampler_id') or None,
            'samp_size':r['_volume'],'samp_size_unit':'mL' if r['_volume'] is not None else None,
            'assay_name':config.get('assay_name') or None,
            'minimumDepthInMeters':float(r['sampling_depth_m']) if r.get('sampling_depth_m') else None,
            'maximumDepthInMeters':float(r['sampling_depth_m']) if r.get('sampling_depth_m') else None,
        }
        for name in selected:
            changes['sampleMetadata'][cell(sample_terms[name],index)]=fields[name]
        missing=[key for key in required_for_this_stage if fields[key] in (None,'')]
        if not r.get('kit_uid'):missing.append('materialSampleID')
        if not fields['assay_name']:missing.append('assay_name')
        if config.get('assay_type')=='targeted':missing.append('detected_notDetected')
        review.append({'sample_event_id':r['sample_event_id'],'missing_fields':'; '.join(missing),
                       'status':'REVIEW_REQUIRED' if missing else 'COMPLETE'})
    patch_workbook(args.template,args.out,changes)
    report_csv(args.out.with_suffix('.review.csv'),['sample_event_id','missing_fields','status'],review)
    missing_project=[x for x in ('recordedBy','project_contact','project_id','assay_type','assay_name') if not config.get(x)]
    print(f'FAIRe draft: {len(rows)} sample rows; unfilled project fields: {", ".join(missing_project) or "none"}')
    print('Laboratory, sequencing, assay detection, and taxon data require later inputs/review.')


if __name__=='__main__':
    try:main()
    except InputError as exc:raise SystemExit(f'Input error: {exc}')
