#!/usr/bin/env python3
"""Add Wilderlab assay and taxon results to an existing FAIRe workbook.

Both taxon sheets mirror processed provider assignments. A separate long CSV
retains per-sample counts because the fixed taxon sheets have no count keys.
"""
import argparse
import hashlib
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from glombo_common import (InputError, cell, patch_workbook, report_csv,
                           sheet_headers)


TAXON_FIELDS={'seq_id','dna_sequence','kingdom','phylum','class','order',
              'family','genus','scientificName','taxonRank','taxonID',
              'verbatimIdentification','identificationRemarks'}
AGGREGATED_FIELDS=['ScientificName','Rank','TaxID','CommonName','Group',
                   'Domain','Kingdom','Phylum','Class','Order','Family','Genus']


def uid(value):
    if value in (None,''):return ''
    if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value).strip()


def label(value):
    return str(value).strip() if value is not None else ''


def taxon_key(name,rank,taxid):
    return uid(taxid),label(name),label(rank)


def read_count(value,kit,row_number):
    if value in ('',None):return 0
    if isinstance(value,bool):raise InputError(f'Invalid count for UID {kit} in full row {row_number}')
    try:number=float(value)
    except (ValueError,TypeError) as exc:
        raise InputError(f'Invalid count for UID {kit} in full row {row_number}') from exc
    if not math.isfinite(number) or number<0 or not number.is_integer():
        raise InputError(f'Expected nonnegative integer count for UID {kit} in full row {row_number}')
    return int(number)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--faire',type=Path,required=True,help='Already generated FAIRe workbook')
    p.add_argument('--results',type=Path,required=True)
    p.add_argument('--out',type=Path,required=True)
    p.add_argument('--allow-extra-results',action='store_true',
                   help='Explicitly permit other UIDs in the same Wilderlab batch')
    p.add_argument('--accept-uid-only',action='store_true',
                   help='After independent review, permit ClientSampleID mismatch')
    p.add_argument('--allow-mock',action='store_true',help='Only for explicit test fixtures')
    args=p.parse_args()
    fw=load_workbook(args.faire,read_only=True,data_only=True)
    sample_header=sheet_headers(args.faire,'sampleMetadata',3)
    fh={key:index-1 for key,index in sample_header.items()}
    required={'samp_name','materialSampleID','assay_name'}
    if not required.issubset(fh):
        raise InputError(f'FAIRe sampleMetadata lacks {sorted(required-set(fh))}')
    by_uid={}
    for row_number,row in enumerate(fw['sampleMetadata'].iter_rows(min_row=4,values_only=True),4):
        sample=label(row[fh['samp_name']])
        kit=uid(row[fh['materialSampleID']])
        if sample and not kit:raise InputError(f'FAIRe sample {sample} has no physical Kit/UID')
        if kit:
            if kit in by_uid:raise InputError(f'Duplicate Kit/UID in FAIRe: {kit}')
            by_uid[kit]=(sample,row_number,label(row[fh['assay_name']]))
    if not by_uid:raise InputError('No samples with physical Kit/UID in FAIRe workbook')
    if any(sample.startswith('MOCK_') for sample,_,_ in by_uid.values()) and not args.allow_mock:
        raise InputError('Mock FAIRe records require --allow-mock')
    for sheet in ('taxaRaw','taxaFinal'):
        if fw[sheet].max_row>3:
            raise InputError(f'{sheet} already contains records; start from the FAIRe sampler workbook')
    project_assay=None
    for row_number,row in enumerate(fw['projectMetadata'].values,1):
        if row[2]=='assay_name':
            project_assay=(row_number,row[3]);break
    if project_assay is None:raise InputError('FAIRe projectMetadata lacks assay_name')
    fw.close()
    results=load_workbook(args.results,read_only=True,data_only=True)
    if not {'metadata','full','aggregated'}.issubset(results.sheetnames):
        raise InputError('Wilderlab result workbook must have metadata, full, aggregated sheets')
    m=results['metadata']
    headers=next(m.iter_rows(min_row=21,max_row=21,values_only=True))
    expected=['UID','ClientSampleID']
    if list(headers[:2])!=expected or 'Assays' not in headers:
        raise InputError('Wilderlab metadata row 21 fields changed')
    assay_index=headers.index('Assays')
    reported={}
    for r in m.iter_rows(min_row=22,values_only=True):
        key=uid(r[0])
        if not key:continue
        if key in reported:raise InputError(f'Duplicate UID in Wilderlab metadata: {key}')
        reported[key]=(label(r[1]),r[assay_index])
    missing=set(by_uid)-set(reported)
    extra=set(reported)-set(by_uid)
    if missing:raise InputError(f'FAIRe Kit/UID absent from Wilderlab results: {sorted(missing)[:10]}')
    if extra and not args.allow_extra_results:
        raise InputError(f'Wilderlab batch has {len(extra)} extra UIDs; use --allow-extra-results only after confirming job scope')
    mismatches=[]
    assays={}
    for key,(sample,_,_) in by_uid.items():
        client_sample,assay_text=reported[key]
        if client_sample!=sample and not args.accept_uid_only:
            raise InputError(f'UID {key}: ClientSampleID {client_sample!r} does not match FAIRe samp_name {sample!r}')
        if client_sample!=sample:mismatches.append(key)
        codes=[part.strip() for part in label(assay_text).split(',') if part.strip()]
        if not codes:raise InputError(f'Wilderlab metadata.Assays is blank for UID {key}')
        if len(codes)!=len(set(codes)):
            raise InputError(f'Duplicate assay code in Wilderlab metadata for UID {key}')
        assays[key]=codes
    ws=results['full']
    h=next(ws.values)
    if list(h[:7])!=['Sequence','Target','ScientificName','Rank','TaxID','CommonName','Group']:
        raise InputError('Wilderlab full sheet fields changed')
    column_uids=[uid(x) for x in h[7:]]
    if len(column_uids)!=len(set(column_uids)) or set(column_uids)!=set(reported):
        raise InputError('Wilderlab full sample columns differ from metadata UIDs')
    aggregated=results['aggregated']
    agg=next(aggregated.values)
    if list(agg[:12])!=AGGREGATED_FIELDS:
        raise InputError('Wilderlab aggregated taxonomic columns changed')
    aggregated_uids=[uid(x) for x in agg[12:]]
    if len(aggregated_uids)!=len(set(aggregated_uids)) or set(aggregated_uids)!=set(reported):
        raise InputError('Wilderlab aggregated sample columns differ from metadata UIDs')
    lineage={}
    for result_row in aggregated.iter_rows(min_row=2,values_only=True):
        if not label(result_row[0]):continue
        key=taxon_key(*result_row[:3])
        fields=tuple(label(value) for value in result_row[3:12])
        if key in lineage and lineage[key]!=fields:
            raise InputError(f'Conflicting aggregated taxonomy for {key}')
        lineage[key]=fields
    selected=[(7+i,key) for i,key in enumerate(column_uids) if key in by_uid]
    taxon_headers={sheet:sheet_headers(args.faire,sheet,3) for sheet in ('taxaRaw','taxaFinal')}
    for sheet,columns in taxon_headers.items():
        if not TAXON_FIELDS.issubset(columns):
            raise InputError(f'FAIRe {sheet} layout has changed')
    edits={'sampleMetadata':{},'taxaRaw':{},'taxaFinal':{}}
    all_assays=list(dict.fromkeys(code for key in by_uid for code in assays[key]))
    if not label(project_assay[1]):
        edits['projectMetadata']={cell(4,project_assay[0]):','.join(all_assays)}
    for key,(_,row_number,_) in by_uid.items():
        edits['sampleMetadata'][cell(sample_header['assay_name'],row_number)]=','.join(assays[key])
    asvs={}
    counts=[]
    summary=defaultdict(lambda:{'asvs':0,'reads':0,'targets':set()})
    for full_row_number,result_row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
        sequence,target,name,rank,taxid=result_row[:5]
        if not sequence or not target:continue
        sequence,target=label(sequence),label(target)
        seq_id='WL_'+hashlib.sha256((target+'|'+sequence).encode()).hexdigest()[:20]
        for index,key in selected:
            amount=read_count(result_row[index],key,full_row_number)
            if amount==0:continue
            if target not in assays[key]:
                raise InputError(f'full row {full_row_number}: target {target} not in metadata.Assays for UID {key}')
            taxon=taxon_key(name,rank,taxid)
            if taxon not in lineage:
                raise InputError(f'full row {full_row_number}: no aggregated taxonomy for {taxon}')
            assignment=(sequence,target,taxon,lineage[taxon],label(result_row[5]),label(result_row[6]))
            if seq_id in asvs and asvs[seq_id]!=assignment:
                raise InputError(f'Sequence ID collision or conflicting assignment: {seq_id}')
            asvs[seq_id]=assignment
            counts.append({'kit_uid':key,'sample_event_id':by_uid[key][0],
                           'seq_id':seq_id,'target_code':target,'sequence_count':amount,
                           'reported_scientific_name':label(name),
                           'reported_taxid':uid(taxid)})
            summary[key]['asvs']+=1;summary[key]['reads']+=amount
            summary[key]['targets'].add(target)
    # Both fixed taxon sheets describe barcode assignments. The processed report
    # is the only classification supplied; mark the mirrored raw rows accordingly.
    for row_number,(seq_id,assignment) in enumerate(sorted(asvs.items()),4):
        sequence,target,taxon,fields,full_common,full_group=assignment
        taxid,name,rank=taxon
        common,group,domain,kingdom,phylum,taxon_class,order,family,genus=fields
        if full_common and common and full_common!=common:
            raise InputError(f'CommonName differs between full and aggregated for {seq_id}')
        if full_group and group and full_group!=group:
            raise InputError(f'Group differs between full and aggregated for {seq_id}')
        details=[f'target code {target}','Wilderlab processed report (full + aggregated)']
        for field,value in (('Domain',domain),('CommonName',common or full_common),('Group',group or full_group)):
            if value:details.append(f'{field}: {value}')
        details.append('Raw pre-filter assignments and taxon ID database not supplied; review pending')
        vals={'seq_id':seq_id,'dna_sequence':sequence,
              'kingdom':kingdom or None,'phylum':phylum or None,
              'class':taxon_class or None,'order':order or None,
              'family':family or None,'genus':genus or None,
              'scientificName':name or None,'taxonRank':rank or None,
              'taxonID':taxid or None,
              'verbatimIdentification':name or None,
              'identificationRemarks':'; '.join(details)+'.'}
        for sheet in ('taxaRaw','taxaFinal'):
            for field,value in vals.items():
                edits[sheet][cell(taxon_headers[sheet][field],row_number)]=value
    results.close()
    patch_workbook(args.faire,args.out,edits)
    report_csv(args.out.with_suffix('.counts_long.csv'),
               ['kit_uid','sample_event_id','seq_id','target_code','sequence_count','reported_scientific_name','reported_taxid'],counts)
    report_csv(args.out.with_suffix('.results_review.csv'),
               ['kit_uid','sample_event_id','wilderlab_client_sample_id','reference_match',
                'reported_assays','previous_sample_assay_name',
                'positive_barcodes','total_reported_sequence_count','target_codes'],
               [{'kit_uid':key,'sample_event_id':sample,'positive_barcodes':summary[key]['asvs'],
                 'wilderlab_client_sample_id':reported[key][0],
                 'reference_match':reported[key][0]==sample,
                 'reported_assays':','.join(assays[key]),
                 'previous_sample_assay_name':old_assay,
                 'total_reported_sequence_count':summary[key]['reads'],
                 'target_codes':';'.join(sorted(summary[key]['targets']))}
                for key,(sample,_,old_assay) in by_uid.items()])
    print(f'Joined {len(by_uid)} UIDs; {len(asvs)} barcode assignments in each taxon sheet; {len(counts)} positive sample/barcode observations')
    if project_assay[1]:
        print('Existing projectMetadata.assay_name preserved; per-sample assays updated from results.')
    if mismatches:print(f'WARNING: {len(mismatches)} ClientSampleID mismatches accepted by explicit UID-only override')
    print('Sequence counts are in the long CSV; fixed FAIRe taxon sheets have no sample/count columns.')


if __name__=='__main__':
    try:main()
    except InputError as exc:raise SystemExit(f'Input error: {exc}')
