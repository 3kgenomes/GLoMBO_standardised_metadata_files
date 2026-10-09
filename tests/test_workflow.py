"""Checks the synthetic three-stage workflow against supplied blank templates."""
import csv
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from glombo_common import InputError, read_samples  # noqa: E402


class PipelineSmokeTest(unittest.TestCase):
    def test_mock_requires_explicit_flag(self):
        source = ROOT / "examples" / "edna_sampler_standardised_MOCK.csv"
        with self.assertRaisesRegex(InputError, "--allow-mock"):
            read_samples(source)

    def test_three_stages(self):
        wilderlab = ROOT / "templates" / "Wilderlab_SampleSubmissionTemplate(1).xlsx"
        faire = ROOT / "templates" / "FAIRe_checklist_v1.0.2_FULLtemplate.xlsx"
        if not wilderlab.exists() or not faire.exists():
            self.skipTest("Place both external blank workbooks in templates/")

        def run(*arguments):
            command = [sys.executable, *map(str, arguments)]
            result = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + "\n" + result.stderr)

        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            sampler = ROOT / "examples" / "edna_sampler_standardised_MOCK.csv"
            result_fixture = ROOT / "examples" / "Wilderlab_results_MOCK.xlsx"
            job = ROOT / "examples" / "wilderlab_job_MOCK.json"
            project = ROOT / "examples" / "faire_project_MOCK.json"
            submission = output / "submission.xlsx"
            initial = output / "initial.xlsx"
            final = output / "final.xlsx"

            run(ROOT / "RVI_2_WilderLab.py", "--samples", sampler,
                "--template", wilderlab, "--job", job,
                "--out", submission, "--allow-mock")
            run(ROOT / "RVI_2_FAIRe.py", "--samples", sampler,
                "--template", faire, "--project", project,
                "--out", initial, "--allow-mock")
            run(ROOT / "Wilderlab_2_FAIRe.py", "--faire", initial,
                "--results", result_fixture, "--out", final, "--allow-mock")

            workbook = load_workbook(final, read_only=True, data_only=True)
            header = [cell.value for cell in workbook["sampleMetadata"][3]]
            assay_column = header.index("assay_name")
            values = [row[assay_column] for row in
                      workbook["sampleMetadata"].iter_rows(min_row=4, values_only=True)]
            self.assertEqual(values, ["WV,CI"] * 3)
            self.assertEqual(workbook["projectMetadata"]["D30"].value, "WV,CI")

            for sheet in ("taxaRaw", "taxaFinal"):
                self.assertEqual(workbook[sheet].max_row, 5)
                taxon_header = [cell.value for cell in workbook[sheet][3]]
                names = {row[taxon_header.index("scientificName")] for row in
                         workbook[sheet].iter_rows(min_row=4, values_only=True)}
                self.assertEqual(names, {"Mockus marinus", "Testa oceanica"})
            workbook.close()

            with final.with_suffix(".counts_long.csv").open(newline="", encoding="utf-8") as f:
                counts = list(csv.DictReader(f))
            self.assertEqual(len(counts), 3)
            self.assertEqual(sum(int(row["sequence_count"]) for row in counts), 22)


if __name__ == "__main__":
    unittest.main()
